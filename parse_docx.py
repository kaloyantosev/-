import docx, json

doc = docx.Document(r'C:\Users\Terve\.gemini\antigravity\brain\ec472d32-9e33-4a2d-9829-7b2aa61d38ca\.user_uploaded\media_1791885624797_319f4a5f.docx')

result = {
    'sections': [],
    'paragraphs': []
}

for s_idx, s in enumerate(doc.sections):
    h = '\n'.join([p.text for p in s.header.paragraphs if p.text.strip()])
    f = '\n'.join([p.text for p in s.footer.paragraphs if p.text.strip()])
    result['sections'].append({
        'index': s_idx,
        'header': h,
        'footer': f,
        'page_width': s.page_width.pt,
        'page_height': s.page_height.pt,
        'top_margin': s.top_margin.pt,
        'bottom_margin': s.bottom_margin.pt,
        'left_margin': s.left_margin.pt,
        'right_margin': s.right_margin.pt,
    })

pages = [[]]
for p_idx, p in enumerate(doc.paragraphs):
    has_break = any('w:br' in r._r.xml and 'type="page"' in r._r.xml for r in p.runs) or 'pagebreakbefore' in p._p.xml.lower()
    if has_break:
        pages.append([])
    
    runs_data = []
    for r in p.runs:
        if r.text:
            runs_data.append({
                'text': r.text,
                'bold': bool(r.bold),
                'italic': bool(r.italic),
                'font': r.font.name,
                'size': r.font.size.pt if r.font.size else None
            })
    p_info = {
        'idx': p_idx,
        'text': p.text,
        'align': str(p.alignment),
        'break': has_break,
        'runs': runs_data
    }
    pages[-1].append(p_info)

with open(r'C:\Users\Terve\.gemini\antigravity\scratch\IRIS подпис на документи\docx_pages_summary.json', 'w', encoding='utf-8') as f:
    json.dump({'sections': result['sections'], 'num_pages': len(pages), 'pages': pages}, f, ensure_ascii=False, indent=2)

print(f'Total pages determined by explicit page breaks: {len(pages)}')
for i, page in enumerate(pages):
    print(f'Page {i+1}: {len(page)} paragraphs. First: {repr(page[0]["text"][:40]) if page else "EMPTY"}, Last: {repr(page[-1]["text"][:40]) if page else "EMPTY"}')
