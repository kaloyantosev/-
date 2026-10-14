import docx

doc = docx.Document(r'C:\Users\Terve\.gemini\antigravity\brain\ec472d32-9e33-4a2d-9829-7b2aa61d38ca\.user_uploaded\media_1791885624797_319f4a5f.docx')

with open(r'C:\Users\Terve\.gemini\antigravity\scratch\IRIS подпис на документи\inspect_docx_detail.txt', 'w', encoding='utf-8') as f:
    f.write(f'Paragraphs: {len(doc.paragraphs)}\n')
    f.write(f'Tables: {len(doc.tables)}\n')
    for i, p in enumerate(doc.paragraphs):
        has_break = any('page' in r._r.xml.lower() for r in p.runs) or 'pagebreakbefore' in p._p.xml.lower()
        f.write(f'P{i} [break={has_break}] [align={p.alignment}]: {p.text}\n')
        for j, r in enumerate(p.runs):
            if r.text:
                f.write(f'   run{j}: bold={r.bold}, italic={r.italic}, font={r.font.name}, size={r.font.size}, text={repr(r.text)}\n')

    for t_idx, table in enumerate(doc.tables):
        f.write(f'\n--- TABLE {t_idx} ---\n')
        for r_idx, row in enumerate(table.rows):
            cells_txt = [c.text.strip().replace('\n', ' ') for c in row.cells]
            f.write(f'Row {r_idx}: ' + ' | '.join(cells_txt) + '\n')

print('Finished inspect_docx')
